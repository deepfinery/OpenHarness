"""Build the source-grounded guide as a typeset PDF and editable Word document.

Run with the pinned packages in requirements.txt. All inputs live in this folder;
no service, database, model, external URL, or sibling checkout is read by this build.
"""

from __future__ import annotations

import html
import re
from pathlib import Path

import fitz
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab import rl_config
from reportlab.graphics import renderPDF
from reportlab.graphics.shapes import Drawing, Line, Polygon, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate, CondPageBreak, Frame, KeepTogether, NextPageTemplate, PageBreak, PageTemplate,
    Paragraph, Spacer, Table, TableStyle,
)
from reportlab.platypus.tableofcontents import TableOfContents

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OUT = HERE / "OpenHarness-Agentic-Capabilities"
BUILD = ROOT / "data" / "agentic-guide-build" / "rendered"
BUILD.mkdir(parents=True, exist_ok=True)
TITLE = "OpenHarness"
SUBTITLE = "Agentic Capabilities & Orchestration"
REVISION = "30b7f628261bf5516c3b3e55871baeb3c80b3973"
REVIEW_DATE = "27 September 2026"
NAVY = "162B42"
TEAL = "087F8C"
INK = "24364B"
MUTED = "5C6B79"
PALE = "EDF5F7"
PAGE_W, PAGE_H = 595.276, 841.89
MARGIN = 47
WIDTH = PAGE_W - 2 * MARGIN


def color(value):
    return colors.HexColor("#" + value)


font_dir = Path(rl_config.TTFSearchPath[0])
for directory in rl_config.TTFSearchPath:
    if (Path(directory) / "Vera.ttf").exists():
        font_dir = Path(directory)
        break
for name, filename in [("Guide", "Vera.ttf"), ("GuideBold", "VeraBd.ttf"),
                       ("GuideItalic", "VeraIt.ttf"), ("GuideMono", "Vera.ttf")]:
    pdfmetrics.registerFont(TTFont(name, str(font_dir / filename)))
pdfmetrics.registerFontFamily("Guide", normal="Guide", bold="GuideBold", italic="GuideItalic", boldItalic="GuideBold")


def inline(text):
    """Render only the small, explicit inline Markdown subset used by this guide."""
    text = html.escape(text)
    text = re.sub(r"\[([^\]]+)\]\((https://[^)]+)\)", r'<link href="\2" color="#087F8C">\1</link>', text)
    text = re.sub(r"`([^`]+)`", r'<font color="#087F8C">\1</font>', text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", text)
    return text


def blocks(text):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            yield "h1", line[2:]
        elif line.startswith("## "):
            yield "h2", line[3:]
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)", line)
            yield "figure", match.groups()
        elif line.startswith("| "):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                row = [c.strip() for c in lines[i].strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", c) for c in row):
                    rows.append(row)
                i += 1
            yield "table", rows
            continue
        elif line.startswith("- "):
            yield "bullet", line[2:]
        elif re.match(r"^\d+\. ", line):
            yield "number", line
        else:
            para = [line]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||!\[|- |\d+\. )", lines[i]):
                para.append(lines[i])
                i += 1
            yield "p", " ".join(para)
            continue
        i += 1


def diagram(kind):
    heights = {"architecture": 260, "lifecycle": 230, "agent-loop": 225, "memory": 255}
    d = Drawing(WIDTH, heights[kind])

    def box(x, y, w, h, title, subtitle="", dark=False):
        d.add(Rect(x, y, w, h, rx=7, ry=7, strokeColor=color(TEAL if dark else "C4D8DF"),
                   fillColor=color(NAVY if dark else PALE), strokeWidth=0.8))
        lines = [title] + ([subtitle] if subtitle else [])
        for n, line in enumerate(lines):
            d.add(String(x + w/2, y + h/2 + (5 if subtitle else -3) - n*14,
                         line, fontName="GuideBold" if n == 0 else "Guide",
                         fontSize=9 if n == 0 else 7.3, textAnchor="middle",
                         fillColor=colors.white if dark else color(INK)))

    def arrow(x1, y1, x2, y2):
        d.add(Line(x1, y1, x2, y2, strokeColor=color(TEAL), strokeWidth=1.2))
        if x1 == x2:
            sign = 1 if y2 > y1 else -1
            points = [x2, y2, x2-3, y2-sign*6, x2+3, y2-sign*6]
        else:
            sign = 1 if x2 > x1 else -1
            points = [x2, y2, x2-sign*6, y2-3, x2-sign*6, y2+3]
        d.add(Polygon(points, fillColor=color(TEAL), strokeColor=color(TEAL)))

    def label(x, y, value):
        d.add(String(x, y, value, fontName="Guide", fontSize=7.5, fillColor=color(MUTED), textAnchor="middle"))

    if kind == "architecture":
        box(0, 195, 135, 52, "Studio / clients", "Chat, API, schedules")
        box(180, 195, 145, 52, "API + dispatcher", "Admission and recovery", True)
        box(365, 195, 135, 52, "RabbitMQ", "Durable delivery")
        arrow(135, 221, 180, 221)
        arrow(325, 221, 365, 221)
        box(365, 105, 135, 52, "Runner pool", "Patterns, tools, children", True)
        arrow(432, 195, 432, 157)
        box(0, 105, 135, 52, "MongoDB", "State, leases, notes")
        box(180, 105, 145, 52, "Shared files", "Bodies and continuations")
        arrow(252, 195, 252, 157)
        arrow(180, 210, 67, 157)
        arrow(365, 130, 325, 130)
        box(0, 15, 150, 52, "Vector store", "Knowledge retrieval")
        box(175, 15, 150, 52, "Models / controls", "LLMs, safety, hooks")
        box(350, 15, 150, 52, "MCP / gateway", "External tools and devices")
        arrow(397, 105, 75, 67)
        arrow(415, 105, 250, 67)
        arrow(432, 105, 425, 67)
    elif kind == "lifecycle":
        box(0, 151, 110, 48, "Accepted", "Snapshot + run")
        box(145, 151, 110, 48, "Queued", "Durable outbox")
        box(290, 151, 110, 48, "Running", "Worker lease", True)
        arrow(110, 175, 145, 175)
        arrow(255, 175, 290, 175)
        box(285, 56, 125, 48, "Human wait", "Saved continuation")
        arrow(330, 151, 330, 104)
        arrow(305, 104, 205, 151)
        label(239, 112, "decision requeues")
        box(0, 56, 145, 48, "Recovery decision", "Resume or interrupt")
        arrow(310, 151, 115, 104)
        arrow(100, 104, 170, 151)
        box(432, 110, 68, 85, "End", "See states")
        arrow(400, 175, 432, 175)
        label(250, 17, "Terminal outcomes: succeeded, failed, cancelled or interrupted")
    elif kind == "agent-loop":
        box(0, 154, 140, 52, "Budget + context", "Fit prompt; save progress")
        box(180, 154, 140, 52, "Model call", "Answer or tool request", True)
        box(360, 154, 140, 52, "Answer", "Validate and finish")
        arrow(140, 180, 180, 180)
        arrow(320, 180, 360, 180)
        box(180, 62, 140, 52, "Controls + journal", "Validate / approve / call")
        box(0, 62, 140, 52, "Observation", "Redact, save, reference")
        arrow(250, 154, 250, 114)
        arrow(180, 88, 140, 88)
        arrow(70, 114, 70, 154)
        box(360, 62, 140, 52, "Human pause", "Resume pending position")
        arrow(320, 88, 360, 88)
        label(250, 22, "Limits reserve tool-free synthesis; ambiguous writes stop recovery")
    else:
        box(0, 185, 145, 50, "Active query", "Agents and children", True)
        box(185, 185, 145, 50, "Task notebook", "MongoDB; seven days")
        box(370, 185, 130, 50, "Promotion", "Selected evidence")
        arrow(145, 210, 185, 210)
        arrow(330, 210, 370, 210)
        box(0, 95, 145, 50, "Terminal run", "Saved outcome")
        box(185, 95, 145, 50, "experiments/", "Automatic; unreviewed")
        box(370, 95, 130, 50, "Notebook", "Files + metadata")
        arrow(72, 185, 72, 145)
        arrow(145, 120, 185, 120)
        arrow(330, 120, 370, 120)
        arrow(435, 185, 435, 145)
        box(0, 5, 145, 50, "Feedback / failure", "Configured learning")
        box(185, 5, 145, 50, "experience/", "Short lesson + provenance")
        box(370, 5, 130, 50, "Later query", "Scoped recall")
        arrow(72, 95, 72, 55)
        arrow(145, 30, 185, 30)
        arrow(330, 30, 370, 30)
        arrow(435, 95, 435, 55)
    return d


STYLES = getSampleStyleSheet()
STYLES.add(ParagraphStyle("BodyGuide", fontName="Guide", fontSize=9.5, leading=14.4,
                          textColor=color(INK), spaceAfter=8, allowWidows=0, allowOrphans=0))
STYLES.add(ParagraphStyle("H1Guide", fontName="GuideBold", fontSize=20, leading=25,
                          textColor=color(NAVY), spaceBefore=16, spaceAfter=15, keepWithNext=True))
STYLES.add(ParagraphStyle("H2Guide", fontName="GuideBold", fontSize=12.1, leading=17,
                          textColor=color(TEAL), spaceBefore=13, spaceAfter=7, keepWithNext=True))
STYLES.add(ParagraphStyle("H2BeforeTable", parent=STYLES["H2Guide"], keepWithNext=False))
STYLES.add(ParagraphStyle("CellGuide", parent=STYLES["BodyGuide"], fontSize=8, leading=11.4, spaceAfter=0))
STYLES.add(ParagraphStyle("CellHead", parent=STYLES["CellGuide"], fontName="GuideBold", textColor=colors.white))
STYLES.add(ParagraphStyle("SourceGuide", parent=STYLES["BodyGuide"], fontSize=8, leading=11.5, textColor=color(MUTED)))
STYLES.add(ParagraphStyle("CaptionGuide", parent=STYLES["SourceGuide"], alignment=TA_CENTER, spaceAfter=12))
STYLES.add(ParagraphStyle("BulletGuide", parent=STYLES["BodyGuide"], leftIndent=12, firstLineIndent=-9))


class GuidePDF(BaseDocTemplate):
    def afterFlowable(self, flowable):
        if isinstance(flowable, Paragraph) and flowable.style.name == "H1Guide":
            title = flowable.getPlainText()
            if not re.match(r"^\d+\.", title):
                return
            key = "section-" + title.split(".", 1)[0]
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(title, key, 0, False)
            self.notify("TOCEntry", (0, title, self.page, key))


def page_frame(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(color("D9E4E8"))
    canvas.line(MARGIN, PAGE_H-40, PAGE_W-MARGIN, PAGE_H-40)
    canvas.setFont("GuideBold", 7)
    canvas.setFillColor(color(TEAL))
    canvas.drawString(MARGIN, PAGE_H-30, "OPENHARNESS  /  AGENTIC CAPABILITIES")
    canvas.setFont("Guide", 7)
    canvas.setFillColor(color(MUTED))
    canvas.drawString(MARGIN, 26, "Source review: " + REVISION[:7] + "  |  " + REVIEW_DATE)
    canvas.drawRightString(PAGE_W-MARGIN, 26, str(doc.page))
    canvas.restoreState()


def cover_frame(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(color(NAVY))
    canvas.rect(0, 0, PAGE_W, PAGE_H, fill=1, stroke=0)
    canvas.setFillColor(color(TEAL))
    canvas.rect(MARGIN, 622, 68, 7, fill=1, stroke=0)
    canvas.setFillColor(colors.white)
    canvas.setFont("GuideBold", 40)
    canvas.drawString(MARGIN, 560, TITLE)
    canvas.setFont("Guide", 26)
    canvas.drawString(MARGIN, 510, "Agentic Capabilities")
    canvas.drawString(MARGIN, 475, "& Orchestration")
    canvas.setFont("Guide", 12)
    for i, line in enumerate(["A detailed implementation guide", "Agent loops · Patterns · Memory · Experiments",
                              "Scalability · Resilience · Human collaboration"]):
        canvas.drawString(MARGIN, 404-i*22, line)
    canvas.setStrokeColor(color("537388"))
    canvas.line(MARGIN, 198, PAGE_W-MARGIN, 198)
    canvas.setFont("GuideBold", 10)
    canvas.drawString(MARGIN, 169, "REPOSITORY REVIEW  /  VERSION 0.2.0")
    canvas.setFont("Guide", 9)
    canvas.drawString(MARGIN, 145, REVIEW_DATE + "  •  Commit " + REVISION[:7])
    canvas.drawString(MARGIN, 121, "deepfinery/OpenHarness")
    canvas.drawString(MARGIN, 75, "Source-grounded descriptions, worked examples and implementation limits")
    canvas.restoreState()


def pdf_table(rows):
    n = len(rows[0])
    ratios = {2: [0.24, 0.76], 3: [0.24, 0.39, 0.37], 4: [0.22, 0.25, 0.26, 0.27],
              5: [0.16, 0.19, 0.22, 0.17, 0.26]}[n]
    cells = [[Paragraph(inline(cell), STYLES["CellHead" if ri == 0 else "CellGuide"])
              for cell in row] for ri, row in enumerate(rows)]
    table = Table(cells, colWidths=[WIDTH*r for r in ratios], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), color(NAVY)),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, color("F2F6F8")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LINEBELOW", (0, 0), (-1, 0), 1, color(TEAL)),
        ("LINEBELOW", (0, 1), (-1, -1), .3, color("DCE5E9")),
    ]))
    table.spaceAfter = 13
    return table


def build_pdf(parsed):
    doc = GuidePDF(str(OUT.with_suffix(".pdf")), pagesize=(PAGE_W, PAGE_H),
                   title=TITLE + ": " + SUBTITLE, author="OpenHarness repository documentation",
                   subject="Agent loops, orchestration patterns, memory, experiments, scalability and resilience")
    doc.addPageTemplates([
        PageTemplate(id="Cover", frames=[Frame(MARGIN, 50, WIDTH, PAGE_H-100)], onPage=cover_frame),
        PageTemplate(id="Body", frames=[Frame(MARGIN, 48, WIDTH, PAGE_H-103,
                                               leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)],
                     onPage=page_frame),
    ])
    toc = TableOfContents()
    toc.levelStyles = [ParagraphStyle("TOCGuide", fontName="Guide", fontSize=9, leading=14,
                                     textColor=color(INK), spaceBefore=0, spaceAfter=0)]
    story = [Spacer(1, 1), NextPageTemplate("Body"), PageBreak(),
             Paragraph("Contents", STYLES["H1Guide"]),
             Paragraph("A reference for engineers, architects and operators. Chapters 4–8 explain agent execution; "
                       "chapters 9–11 cover memory and experiments; chapters 15–17 cover recovery and scale.", STYLES["BodyGuide"]),
             Spacer(1, 10), toc]
    figure_n = 0
    for index, (kind, value) in enumerate(parsed):
        if kind == "h1":
            boundary = PageBreak() if value.startswith("1. ") else CondPageBreak(250)
            story.extend([boundary, Paragraph(inline(value), STYLES["H1Guide"])])
        elif kind == "h2":
            before_table = index + 1 < len(parsed) and parsed[index + 1][0] == "table"
            if before_table:
                story.append(CondPageBreak(135))
            story.append(Paragraph(inline(value), STYLES["H2BeforeTable" if before_table else "H2Guide"]))
        elif kind == "table":
            story.append(pdf_table(value))
        elif kind == "figure":
            figure_n += 1
            caption, key = value
            story.append(KeepTogether([diagram(key), Spacer(1, 6),
                                      Paragraph(f"Figure {figure_n}. {caption}", STYLES["CaptionGuide"])]))
        else:
            style = "SourceGuide" if value.startswith("Source basis:") or re.match(r"^\[S\d+\]", value) else "BodyGuide"
            if kind == "bullet":
                value = "• " + value
                style = "BulletGuide"
            story.append(Paragraph(inline(value), STYLES[style]))
    doc.multiBuild(story)


def word_inline(p, text):
    pattern = r"(\[[^\]]+\]\(https://[^)]+\)|`[^`]+`|\*\*[^*]+\*\*)"
    for part in re.split(pattern, text):
        match = re.fullmatch(r"\[([^]]+)\]\((https://[^)]+)\)", part)
        if match:
            rel = p.part.relate_to(match[2], "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), rel)
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            c = OxmlElement("w:color")
            c.set(qn("w:val"), TEAL)
            props.append(c)
            run.append(props)
            content = OxmlElement("w:t")
            content.text = match[1]
            run.append(content)
            link.append(run)
            p._p.append(link)
        else:
            code = part.startswith("`") and part.endswith("`")
            bold = part.startswith("**") and part.endswith("**")
            run = p.add_run(part[1:-1] if code else part[2:-2] if bold else part)
            if code:
                run.font.color.rgb = RGBColor.from_string(TEAL)
            run.bold = bold


def build_word(parsed):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    section.top_margin, section.bottom_margin = Inches(.72), Inches(.65)
    section.left_margin = section.right_margin = Inches(.66)
    section.header_distance = section.footer_distance = Inches(.3)
    section.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, ink in [("Heading 1", 23, NAVY), ("Heading 2", 13, TEAL)]:
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor.from_string(ink)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(10)
    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    header = section.header.paragraphs[0]
    header.add_run("OPENHARNESS  /  AGENTIC CAPABILITIES").font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.add_run("Source: " + REVISION[:7] + "  |  " + REVIEW_DATE + "                                      ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(120)
    r = p.add_run(TITLE)
    r.bold, r.font.size = True, Pt(44)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph("Agentic Capabilities\n& Orchestration")
    for r in p.runs:
        r.font.size = Pt(29)
        r.font.color.rgb = RGBColor.from_string(TEAL)
    doc.add_paragraph("A detailed implementation guide\nAgent loops · Patterns · Memory · Experiments\nScalability · Resilience · Human collaboration")
    p = doc.add_paragraph("Repository version 0.2.0\n" + REVIEW_DATE + "\nReviewed commit: " + REVISION)
    p.paragraph_format.space_before = Pt(65)
    doc.add_paragraph("Source-grounded descriptions, worked examples and implementation limits.")
    doc.add_heading("Contents", 1)
    doc.add_paragraph("Chapter index. The PDF includes page numbers and bookmarks; Word headings support the Navigation pane.")
    for kind, value in parsed:
        if kind == "h1":
            p = doc.add_paragraph(value)
            p.paragraph_format.space_after = Pt(3)
    figure_n = 0
    for kind, value in parsed:
        if kind == "h1":
            doc.add_heading(value, 1)
        elif kind == "h2":
            doc.add_heading(value, 2)
        elif kind == "table":
            table = doc.add_table(rows=0, cols=len(value[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Light Shading Accent 1"
            for ri, row in enumerate(value):
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(4)
                    word_inline(p, text)
                    for run in p.runs:
                        run.font.size = Pt(9)
                        if ri == 0:
                            run.bold = True
                props = table.rows[-1]._tr.get_or_add_trPr()
                props.append(OxmlElement("w:cantSplit"))
                if ri == 0:
                    props.append(OxmlElement("w:tblHeader"))
            doc.add_paragraph()
        elif kind == "figure":
            figure_n += 1
            caption, key = value
            diagram_pdf = BUILD / (key + ".pdf")
            renderPDF.drawToFile(diagram(key), str(diagram_pdf))
            with fitz.open(diagram_pdf) as image_doc:
                image_doc[0].get_pixmap(matrix=fitz.Matrix(2.4, 2.4), alpha=False).save(str(BUILD / (key + ".png")))
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(BUILD / (key + ".png")), width=Inches(6.9))
            p = doc.add_paragraph(f"Figure {figure_n}. {caption}", "Caption")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p = doc.add_paragraph(style="List Bullet" if kind == "bullet" else "Normal")
            word_inline(p, value)
            if value.startswith("Source basis:") or re.match(r"^\[S\d+\]", value):
                for run in p.runs:
                    run.font.size = Pt(9)
    props = doc.core_properties
    props.title = TITLE + ": " + SUBTITLE
    props.subject = "Agent loops, orchestration patterns, memory, experiments, scalability and resilience"
    props.author = "OpenHarness repository documentation"
    props.keywords = "OpenHarness, agents, orchestration, experiments, resilience"
    doc.save(OUT.with_suffix(".docx"))


if __name__ == "__main__":
    source = (HERE / "guide.md").read_text()
    assert "{{SOURCE_REGISTER}}" not in source, "Source register must be populated"
    parsed = list(blocks(source))
    build_pdf(parsed)
    build_word(parsed)
    with fitz.open(OUT.with_suffix(".pdf")) as pdf:
        print(f"Built {len(pdf)} PDF pages, {len(source.split()):,} source words, PDF + DOCX")
